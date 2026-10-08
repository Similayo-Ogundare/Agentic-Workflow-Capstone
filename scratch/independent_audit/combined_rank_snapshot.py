from pathlib import Path
from collections import Counter
from datetime import date,datetime
import calendar
import hashlib
import json
import sys
import re
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

BASE=Path('scratch/independent_audit')
audit_path=BASE/'current_audit.json'
audit=json.loads(audit_path.read_text(encoding='utf-8'))
old_hashes={f['file']:f['sha256'] for f in audit['files']}
months={'January':1,'February':2,'March':3,'April':4}
records=[]
files=[]
rosters=[]
original_paths=sorted(Path('Marvellous Church/1st Quarter').rglob('*.xlsx'))
for path in original_paths:
    month,year=path.parent.name.split()
    day=int(re.search(r'(\d+)\.xlsx$',path.name).group(1))
    service_date=date(int(year),months[month],day)
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    roster=[]
    for row_no,row in enumerate(wb.active.iter_rows(values_only=True),start=1):
        if row_no==1:
            continue
        first,last,tithe,offering=row
        name=f'{first} {last}'
        roster.append(name)
        records.append({'name':name,'month':f'{month} {year}','date':service_date.isoformat(),'tithe':tithe if isinstance(tithe,(int,float)) else 0,'offering':offering if isinstance(offering,(int,float)) else 0,'tithe_paid':int(isinstance(tithe,(int,float)) and tithe>0),'offering_paid':int(isinstance(offering,(int,float)) and offering>0)})
    wb.close()
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    files.append({'file':str(path),'sha256':digest,'matches_prior':digest==old_hashes.get(str(path)),'date':service_date.isoformat(),'members':len(roster)})
    rosters.append(roster)

metrics=lambda rows:{k:sum(r[k] for r in rows) for k in ('tithe','offering','tithe_paid','offering_paid')}
fresh_metrics={n:metrics([r for r in records if r['name']==n]) for n in rosters[0]}
fresh_months={m:metrics([r for r in records if r['month']==m]) for m in audit['monthly']}
assert all(audit['members'][n]['final'][k]==v for n,m in fresh_metrics.items() for k,v in m.items()),'Source metrics differ from prior audit'
assert fresh_months==audit['monthly'],'Monthly totals differ from prior audit'
for member in rosters[0]:
    for month in audit['monthly']:
        monthly_member=metrics([r for r in records if r['name']==member and r['month']==month])
        assert monthly_member==audit['members'][member][month],f'Monthly metrics differ for{member}/{month}'
assert all(roster==audit['roster'] for roster in rosters),'Roster differs'
audit['files']=files
audit['all_fingerprints_match_prior']=all(f['matches_prior'] for f in files)
audit['source_inventory_matches_prior']={str(p) for p in original_paths}==set(old_hashes)
audit['file_count']=len(files)
audit['row_count']=len(records)
audit['roster_consistent']=all(roster==rosters[0] for roster in rosters)
audit['calendars']={}
for month,month_number in months.items():
    observed=sorted({r['date'] for r in records if r['month']==f'{month} 2025'})
    expected=[date(2025,month_number,day).isoformat() for day in range(1,calendar.monthrange(2025,month_number)[1]+1) if date(2025,month_number,day).weekday()==6]
    audit['calendars'][month]={'dates':observed,'expected':expected,'complete':observed==expected}
assert all(c['complete'] for c in audit['calendars'].values()),'Sunday calendars are incomplete'
assert len(rosters[0])==19 and len(set(rosters[0]))==19,'Expected19 unique members'
audit['ranks_combined']={}
audit['exact_triple_ties']={}
changes=[]
for kind in ('tithe','offering'):
    def key(n):
        m=fresh_metrics[n]
        return (m[kind+'_paid'],m[kind],m['tithe']+m['offering'])
    grouped={}
    for n in rosters[0]:
        grouped.setdefault(key(n),[]).append(n)
    audit['exact_triple_ties'][kind]=[{'key':k,'names':names} for k,names in grouped.items() if len(names)>1]
    ordered=sorted(rosters[0],key=lambda n:tuple(-x for x in key(n)))
    audit['ranks_combined'][kind]=[]
    for n in ordered:
        rank=1+sum(key(other)>key(n) for other in rosters[0])
        old_rank=audit['members'][n]['final'][kind+'_rank_shared']
        audit['members'][n]['final'][kind+'_rank_combined']=rank
        audit['ranks_combined'][kind].append({'name':n,'paid':key(n)[0],'amount':key(n)[1],'combined_amount':key(n)[2],'rank':rank})
        if rank!=old_rank:
            changes.append({'name':n,'category':kind,'previous_rank':old_rank,'new_rank':rank,'triple_key':key(n)})
audit['combined_rank_rule']='Descending payment occurrence count, then category amount, then combined tithe+offering amount. Competition rank for any remaining exact triple tie.'
audit_path.write_text(json.dumps(audit,indent=2),encoding='utf-8')

NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
def json_value(v):
    return v.isoformat() if isinstance(v,(date,datetime)) else v

def snapshot(path):
    result={'file':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'parts':{},'sheets':{}}
    with zipfile.ZipFile(path) as z:
        for part_name in sorted(z.namelist()):
            raw=z.read(part_name)
            part={'sha256':hashlib.sha256(raw).hexdigest(),'length':len(raw)}
            if part_name.endswith(('.xml','.rels')):
                part['xml']=raw.decode('utf-8')
            result['parts'][part_name]=part
        workbook=ET.fromstring(z.read('xl/workbook.xml'))
        relations=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        rels={r.attrib['Id']:r.attrib['Target'] for r in relations}
        for sheet in workbook.findall('s:sheets/s:sheet',NS):
            target=rels[sheet.attrib['{'+NS['r']+'}id']]
            part=target.lstrip('/') if target.startswith('/') else 'xl/'+target
            root=ET.fromstring(z.read(part))
            cells={}
            for cell in root.findall('s:sheetData/s:row/s:c',NS):
                cells[cell.attrib['r']]={'attributes':cell.attrib,'xml':ET.tostring(cell,encoding='unicode')}
            other_xml=[ET.tostring(e,encoding='unicode') for e in root if e.tag!='{'+NS['s']+'}sheetData']
            result['sheets'][sheet.attrib['name']]={'part':part,'workbook_attributes':sheet.attrib,'cells':cells,'non_cell_xml':other_xml}
    wb_formula=openpyxl.load_workbook(path,read_only=True,data_only=False)
    wb_values=openpyxl.load_workbook(path,read_only=True,data_only=True)
    for s in wb_formula.worksheets:
        sheet_result=result['sheets'][s.title]
        sheet_result['openpyxl_cells']={}
        values=wb_values[s.title]
        value_rows=list(values.iter_rows(values_only=True))
        for row_no,row in enumerate(s.iter_rows(),start=1):
            for col_no,c in enumerate(row,start=1):
                if c.value is not None:
                    sheet_result['openpyxl_cells'][c.coordinate]={'value':json_value(c.value),'cached_value':json_value(value_rows[row_no-1][col_no-1]),'data_type':c.data_type,'style_id':c._style_id,'number_format':c.number_format}
    wb_formula.close(); wb_values.close()
    return result

snapshots={label:snapshot(Path(f'outputs/{label}_Summary.xlsx')) for label in ('January','February','March','April','Final')}
snapshot_path=BASE/(sys.argv[1] if len(sys.argv)>1 else 'rerun_before_snapshot.json')
snapshot_path.write_text(json.dumps(snapshots,indent=2),encoding='utf-8')
print(json.dumps({'sources':len(files),'unchanged':audit['all_fingerprints_match_prior'],'source_inventory_matches_prior':audit['source_inventory_matches_prior'],'member_records':len(records),'all_calendars_complete':all(c['complete'] for c in audit['calendars'].values()),'ranks_changed_from_original_shared':changes,'exact_triple_ties':audit['exact_triple_ties'],'snapshot_files':list(snapshots),'snapshot_path':str(snapshot_path),'snapshot_size':snapshot_path.stat().st_size},indent=2))
