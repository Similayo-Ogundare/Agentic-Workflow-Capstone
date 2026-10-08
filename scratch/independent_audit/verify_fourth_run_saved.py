from pathlib import Path
from datetime import date,datetime
import hashlib
import json
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

BASE=Path('scratch/independent_audit')
before=json.loads((BASE/'fourth_run_before_snapshot.json').read_text(encoding='utf-8'))
expected=json.loads((BASE/'fourth_run_expected.json').read_text(encoding='utf-8'))
allowed={f'J{row}' for row in range(7,26)}|{'A32'}
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS={'s':S}
issues=[]
checks=0
def canon(node):
    return(node.tag,tuple(sorted(node.attrib.items())),node.text or '',tuple(canon(c) for c in node))
def json_value(v):
    return v.isoformat() if isinstance(v,(date,datetime)) else v
def normalize_relationships(root,part,parts):
    relpart=part if part.endswith('.rels') else str(Path(part).parent/'_rels'/f'{Path(part).name}.rels').replace('\\','/')
    if relpart not in parts:
        return
    relations=ET.fromstring(parts[relpart])
    mapping={n.attrib['Id']:json.dumps({k:v for k,v in n.attrib.items() if k!='Id'},sort_keys=True) for n in relations}
    if part.endswith('.rels'):
        for node in root:
            node.attrib['Id']=mapping[node.attrib['Id']]
    else:
        for node in root.iter():
            for key in node.attrib:
                if key.startswith('{'+R+'}') and node.attrib[key] in mapping:
                    node.attrib[key]=mapping[node.attrib[key]]

monthly=[]
for label in ('January','February','March','April'):
    unchanged=hashlib.sha256(Path(before[label]['file']).read_bytes()).hexdigest()==before[label]['sha256']
    monthly.append({'month':label,'byte_identical':unchanged})
    if not unchanged:
        issues.append({'type':'monthly file changed','month':label})

old=before['Final']
path=Path(old['file'])
wf=openpyxl.load_workbook(path,read_only=True,data_only=False)
wv=openpyxl.load_workbook(path,read_only=True,data_only=True)
changes=[]
rank_results=[]
if wf.sheetnames!=list(old['sheets']):
    issues.append({'type':'sheet name inventory changed'})
for sheet in wf.worksheets:
    actual={}
    cached=list(wv[sheet.title].iter_rows(values_only=True))
    for row_num,row in enumerate(sheet.iter_rows(),start=1):
        for col_num,cell in enumerate(row,start=1):
            if cell.value is not None:
                actual[cell.coordinate]={'value':json_value(cell.value),'cached_value':json_value(cached[row_num-1][col_num-1]),'data_type':cell.data_type,'style_id':cell._style_id,'number_format':cell.number_format}
    prior=old['sheets'][sheet.title]['openpyxl_cells']
    for coord in set(actual)|set(prior):
        checks+=1
        if actual.get(coord)!=prior.get(coord):
            permitted=sheet.title=='Summary' and coord in allowed
            change={'sheet':sheet.title,'cell':coord,'permitted':permitted,'before':prior.get(coord),'after':actual.get(coord)}
            changes.append(change)
            if not permitted:
                issues.append({'type':'unrelated cell changed',**change})
            elif actual.get(coord) is not None and prior.get(coord) is not None:
                for key in ('style_id','number_format','data_type'):
                    if actual[coord][key]!=prior[coord][key]:
                        issues.append({'type':'authorized changed cell style/type altered','cell':coord,'key':key})
    if sheet.title=='Summary':
        for member in expected['members']:
            row=member['row']
            formula=f'=1+COUNTIF($I$7:$I$25,"<"&I{row})+SUMPRODUCT(($I$7:$I$25=I{row})*(($C$7:$C$25+$D$7:$D$25)>(C{row}+D{row})))'
            cell=actual['J'+str(row)]
            checks+=2
            if cell['value']!=formula:
                issues.append({'type':'overall formula mismatch','cell':'J'+str(row),'actual':cell['value'],'expected':formula})
            if cell['cached_value']!=member['overall_rank']:
                issues.append({'type':'overall cached value mismatch','name':member['name'],'actual':cell['cached_value'],'expected':member['overall_rank']})
            rank_results.append({'name':member['name'],'overall_rank':cell['cached_value']})
wf.close();wv.close()

with zipfile.ZipFile(path) as z:
    new_parts={part:z.read(part) for part in z.namelist()}
old_parts={part:data['xml'].encode('utf-8') for part,data in old['parts'].items() if 'xml' in data}
if set(new_parts)!=set(old['parts']):
    issues.append({'type':'ZIP part inventory changed'})
serialization_only=[]
changed_parts=[]
for part,data in old['parts'].items():
    if part not in new_parts:
        continue
    if hashlib.sha256(new_parts[part]).hexdigest()==data['sha256']:
        continue
    changed_parts.append(part)
    if 'xml' not in data:
        issues.append({'type':'binary native part changed','part':part});continue
    a=ET.fromstring(data['xml']);b=ET.fromstring(new_parts[part])
    normalize_relationships(a,part,old_parts);normalize_relationships(b,part,new_parts)
    if canon(a)==canon(b):
        serialization_only.append(part);continue
    if part==old['sheets']['Summary']['part']:
        old_cells={c.attrib['r']:c for c in a.findall('s:sheetData/s:row/s:c',NS)}
        new_cells={c.attrib['r']:c for c in b.findall('s:sheetData/s:row/s:c',NS)}
        for coord in allowed:
            if old_cells.get(coord) is not None and new_cells.get(coord) is not None and old_cells[coord].attrib!=new_cells[coord].attrib:
                issues.append({'type':'authorized cell XML attributes/style changed','cell':coord})
        for root in (a,b):
            for row in root.findall('s:sheetData/s:row',NS):
                for cell in list(row):
                    if cell.attrib.get('r') in allowed:
                        row.remove(cell)
        if canon(a)==canon(b):
            continue
    issues.append({'type':'unrelated native XML changed','part':part})

source_changes=[f['file'] for f in expected['source_files'] if hashlib.sha256(Path(f['file']).read_bytes()).hexdigest()!=f['sha256']]
current_inventory=sorted(str(p.resolve()) for p in Path('outputs').rglob('*.xlsx') if not p.name.startswith('~$'))
inventory_unchanged=current_inventory==expected['output_workbook_inventory']
issues.extend({'type':'source changed','file':p} for p in source_changes)
if not inventory_unchanged:
    issues.append({'type':'output workbook inventory changed'})
unique=sorted(r['overall_rank'] for r in rank_results)==list(range(1,20))
if not unique:
    issues.append({'type':'overall ranks not unique1–19'})
report={'pass':not issues,'checks':checks,'issues':issues,'monthly_files':monthly,'source_hash_changes':source_changes,'output_workbook_inventory_unchanged':inventory_unchanged,'unique_overall_1_to_19':unique,'cell_changes':changes,'changed_parts':changed_parts,'serialization_only_parts':serialization_only,'saved_ranks':rank_results}
(BASE/'fourth_run_saved_qa.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ('cell_changes','saved_ranks')},indent=2))
print(json.dumps({'changed_cells':[(c['sheet'],c['cell']) for c in changes],'numerical_rank_changes':[{'cell':c['cell'],'before':c['before']['cached_value'],'after':c['after']['cached_value']} for c in changes if c['cell'].startswith('J') and c['before']['cached_value']!=c['after']['cached_value']]},indent=2))
