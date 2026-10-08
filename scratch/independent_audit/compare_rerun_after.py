from pathlib import Path
from datetime import date,datetime
import copy
import hashlib
import json
import sys
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

BASE=Path('scratch/independent_audit')
before=json.loads((BASE/'rerun_before_snapshot.json').read_text(encoding='utf-8'))
audit=json.loads((BASE/'current_audit.json').read_text(encoding='utf-8'))
require_unchanged='--unchanged' in sys.argv
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS={'s':S}
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
allowed={f'{col}{r}' for col in ('F','H') for r in range(7,26)}|{'A29'}
def canonical(node):
    return (node.tag,tuple(sorted(node.attrib.items())),node.text or '',tuple(canonical(c) for c in node))
def norm_xml(xml):
    return canonical(ET.fromstring(xml))
def value(v):
    return v.isoformat() if isinstance(v,(date,datetime)) else v
def rel_maps(parts):
    maps={}
    for part,xml in parts.items():
        if not part.endswith('.rels'):
            continue
        rels=ET.fromstring(xml)
        mapping={r.attrib['Id']:json.dumps({k:v for k,v in r.attrib.items() if k!='Id'},sort_keys=True) for r in rels}
        maps[part]=mapping
        owner=part.replace('/_rels/','/')[:-5] if part!='_rels/.rels' else ''
        maps[owner]=mapping
    return maps
def normalize_rel_ids(root,part,maps):
    mapping=maps.get(part,{})
    if part.endswith('.rels'):
        for relation in root:
            if 'Id' in relation.attrib:
                relation.attrib['Id']=mapping[relation.attrib['Id']]
    else:
        for element in root.iter():
            for key in list(element.attrib):
                if key.startswith('{'+R+'}'):
                    rel_id=element.attrib[key]
                    if rel_id in mapping:
                        element.attrib[key]=mapping[rel_id]
results={}
all_issues=[]
for label,old in before.items():
    path=Path(old['file'])
    issues=[]
    changes=[]
    wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
    wb_values=openpyxl.load_workbook(path,read_only=True,data_only=True)
    new_cells={}
    if wb.sheetnames!=list(old['sheets']):
        issues.append({'type':'sheet names','before':list(old['sheets']),'after':wb.sheetnames})
    for sheet in wb.worksheets:
        value_rows=list(wb_values[sheet.title].iter_rows(values_only=True))
        cells={}
        for row_index,row in enumerate(sheet.iter_rows(),start=1):
            for col_index,c in enumerate(row,start=1):
                if c.value is not None:
                    cells[c.coordinate]={'value':value(c.value),'cached_value':value(value_rows[row_index-1][col_index-1]),'data_type':c.data_type,'style_id':c._style_id,'number_format':c.number_format}
        new_cells[sheet.title]=cells
        prior=old['sheets'][sheet.title]['openpyxl_cells']
        for coord in sorted(set(cells)|set(prior)):
            if cells.get(coord)!=prior.get(coord):
                permitted=not require_unchanged and label=='Final' and sheet.title=='Summary' and coord in allowed
                change={'sheet':sheet.title,'cell':coord,'before':prior.get(coord),'after':cells.get(coord),'permitted':permitted}
                changes.append(change)
                if not permitted:
                    issues.append({'type':'unrelated cell change',**change})
                elif prior.get(coord) is not None and cells.get(coord) is not None:
                    for key in ('style_id','number_format'):
                        if prior[coord][key]!=cells[coord][key]:
                            issues.append({'type':'allowed cell style changed','cell':coord,'key':key,'before':prior[coord][key],'after':cells[coord][key]})
    if label=='Final':
        final_cells=new_cells['Summary']
        for row in range(7,26):
            name=f"{final_cells['A'+str(row)]['cached_value']} {final_cells['B'+str(row)]['cached_value']}"
            for col,key in [('F','tithe_rank_combined'),('H','offering_rank_combined')]:
                actual=final_cells[col+str(row)]['cached_value']
                expected=audit['members'][name]['final'][key]
                if actual!=expected:
                    issues.append({'type':'rank result mismatch','cell':col+str(row),'name':name,'actual':actual,'expected':expected})
    wb.close();wb_values.close()
    changed_parts=[]
    relationship_id_only_parts=[]
    semantic_part_issues=[]
    with zipfile.ZipFile(path) as z:
        new_names=set(z.namelist())
        old_names=set(old['parts'])
        old_maps=rel_maps({part:data['xml'] for part,data in old['parts'].items() if 'xml' in data})
        new_maps=rel_maps({part:z.read(part).decode('utf-8') for part in new_names if part.endswith('.rels')})
        if new_names!=old_names:
            issues.append({'type':'zip part inventory','removed':sorted(old_names-new_names),'added':sorted(new_names-old_names)})
        for part_name in sorted(new_names&old_names):
            raw=z.read(part_name)
            previous=old['parts'][part_name]
            digest=hashlib.sha256(raw).hexdigest()
            if digest==previous['sha256']:
                continue
            changed_parts.append(part_name)
            if 'xml' not in previous:
                semantic_part_issues.append({'part':part_name,'reason':'binary part changed'})
                continue
            old_root=ET.fromstring(previous['xml'])
            new_root=ET.fromstring(raw)
            if canonical(old_root)==canonical(new_root):
                continue
            normalize_rel_ids(old_root,part_name,old_maps)
            normalize_rel_ids(new_root,part_name,new_maps)
            if canonical(old_root)==canonical(new_root):
                relationship_id_only_parts.append(part_name)
                continue
            if not require_unchanged and label=='Final' and part_name==old['sheets']['Summary']['part']:
                for root in (old_root,new_root):
                    for row in root.findall('s:sheetData/s:row',NS):
                        for cell in list(row):
                            if cell.attrib.get('r') in allowed:
                                row.remove(cell)
                if canonical(old_root)==canonical(new_root):
                    continue
            if not require_unchanged and label=='Final' and part_name=='xl/sharedStrings.xml':
                prior_text=old['sheets']['Summary']['openpyxl_cells']['A29']['value']
                current_text=new_cells['Summary']['A29']['value']
                for root in (old_root,new_root):
                    for t in root.iter('{'+S+'}t'):
                        if t.text in (prior_text,current_text):
                            t.text='__AUTHORIZED_A29_EXPLANATION__'
                if canonical(old_root)==canonical(new_root):
                    continue
            semantic_part_issues.append({'part':part_name,'reason':'XML content changed outside expected rank cells and A29'})
    issues.extend({'type':'unrelated native part change',**p} for p in semantic_part_issues)
    result={'sha256_before':old['sha256'],'sha256_after':hashlib.sha256(path.read_bytes()).hexdigest(),'cell_changes':changes,'changed_parts':changed_parts,'relationship_id_only_parts':relationship_id_only_parts,'issues':issues}
    results[label]=result
    all_issues.extend({'workbook':label,**issue} for issue in issues)
source_changes=[f['file'] for f in audit['files'] if hashlib.sha256(Path(f['file']).read_bytes()).hexdigest()!=f['sha256']]
report={'workbooks':results,'issues':all_issues,'source_hash_changes':source_changes,'require_semantic_unchanged':require_unchanged,'all_allowed_changes_only':not all_issues and not source_changes}
(BASE/'rerun_after_comparison.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps({'all_allowed_changes_only':report['all_allowed_changes_only'],'source_hash_changes':source_changes,'issues':all_issues,'workbooks':{k:{'changed_parts':v['changed_parts'],'relationship_id_only_parts':v['relationship_id_only_parts'],'changed_cells':[(c['sheet'],c['cell']) for c in v['cell_changes']],'issue_count':len(v['issues'])} for k,v in results.items()}},indent=2))
