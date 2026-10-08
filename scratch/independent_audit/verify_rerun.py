from pathlib import Path
from datetime import date,datetime
import contextlib
import hashlib
import io
import json
import runpy
import openpyxl

audit_path=Path('scratch/independent_audit/current_audit.json')
previous=json.loads(audit_path.read_text(encoding='utf-8'))
with contextlib.redirect_stdout(io.StringIO()):
    fresh=runpy.run_path('scratch/independent_audit/current_audit.py')
audit=json.loads(audit_path.read_text(encoding='utf-8'))
prior_hashes={entry['file']:entry['sha256'] for entry in previous['files']}
hash_changes=[entry['file'] for entry in audit['files'] if prior_hashes.get(entry['file'])!=entry['sha256']]
checks=0
discrepancies=[]
output_snapshots=[]
def compare(actual,expected,location):
    global checks
    checks+=1
    if actual!=expected:
        discrepancies.append({'location':location,'actual':actual,'expected':expected})

for label in ('January','February','March','April','Final'):
    path=Path(f'outputs/{label}_Summary.xlsx')
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    sheet=wb['Summary']
    rows=list(sheet.iter_rows(values_only=True))
    header_index=next(i for i,r in enumerate(rows) if r[:2]==('First Name','Last Name'))
    member_rows=rows[header_index+1:header_index+20]
    member_names=[f'{r[0]} {r[1]}' for r in member_rows]
    compare(member_names,audit['roster'],f'{path}/Summary/member roster')
    for row_index,row in enumerate(member_rows,start=header_index+2):
        name=f'{row[0]} {row[1]}'
        if name not in audit['members']:
            continue
        month_key=f'{label} 2025' if label!='Final' else 'final'
        values=audit['members'][name][month_key]
        fields=('tithe','offering','tithe_paid','offering_paid') if label!='Final' else ('tithe','offering','tithe_paid','tithe_rank_shared','offering_paid','offering_rank_shared')
        for col_index,metric in enumerate(fields,start=2):
            compare(row[col_index],values[metric],f'{path}/Summary/R{row_index}C{col_index+1} ({name}/{metric})')
    totals_row=next(r for r in rows if r[0]=='Total')
    expected=audit['monthly'][f'{label} 2025'] if label!='Final' else audit['final']
    if label!='Final':
        for c,k in enumerate(('tithe','offering','tithe_paid','offering_paid'),start=2):
            compare(totals_row[c],expected[k],f'{path}/Summary/Church total/{k}')
    else:
        for c,k in ((2,'tithe'),(3,'offering'),(4,'tithe_paid'),(6,'offering_paid')):
            compare(totals_row[c],expected[k],f'{path}/Summary/Church total/{k}')
    if label!='Final':
        source_rows=list(wb['Source records'].iter_rows(values_only=True))[4:]
        source_records=sorted([r for r in fresh['records'] if r['month']==label],key=lambda r:(r['date'],r['row']))
        compare(len(source_rows),len(source_records),f'{path}/Source records/row count')
        for i,(row,record) in enumerate(zip(source_rows,source_records),start=5):
            raw=lambda k:'N/A' if record[k+'_status']=='N/A' else None if record[k+'_status']=='blank' else record[k]
            expected_row=(datetime.combine(date.fromisoformat(record['date']),datetime.min.time()),record['first'],record['last'],raw('tithe'),raw('offering'),record['tithe_paid'],record['offering_paid'],Path(record['file']).name,record['row'])
            compare(row,expected_row,f'{path}/Source records/row{i}')
    else:
        monthly_rows=list(wb['Monthly totals'].iter_rows(values_only=True))[4:]
        expected_rows=[]
        for month_number,label_month in enumerate(('January','February','March','April'),start=1):
            for name in audit['roster']:
                first,last=name.split(' ',1)
                metrics=audit['members'][name][f'{label_month} 2025']
                expected_rows.append((datetime(2025,month_number,1),first,last,metrics['tithe'],metrics['offering'],metrics['tithe_paid'],metrics['offering_paid'],f'{label_month}_Summary.xlsx'))
        compare(len(monthly_rows),len(expected_rows),f'{path}/Monthly totals/row count')
        for i,(actual,expected_row) in enumerate(zip(monthly_rows,expected_rows),start=5):
            compare(actual,expected_row,f'{path}/Monthly totals/row{i}')
    for s in wb.worksheets:
        for row in s.iter_rows():
            for cell in row:
                if cell.data_type=='e':
                    discrepancies.append({'location':f'{path}/{s.title}/{cell.coordinate}','error':cell.value})
    wb.close()
    output_snapshots.append({'file':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'members':len(member_rows),'sheets':['Summary','Monthly totals'] if label=='Final' else ['Summary','Source records']})

result={'source_files':len(audit['files']),'source_hash_changes':hash_changes,'source_inventory_matches_prior':{e['file'] for e in previous['files']}=={e['file'] for e in audit['files']},'roster_consistent':audit['roster_consistent'],'calendars_complete':all(entry['complete'] for entry in audit['calendars'].values()),'checks':checks,'output_discrepancies':discrepancies,'output_snapshots':output_snapshots,'final':audit['final'],'monthly':audit['monthly'],'exact_ties':audit['exact_ties'],'existing_output_tie_mode':'Shared competition ranks','note':'Output values match fresh source-derived amounts/counts and prior competition ranks when no discrepancies. No full baseline output hash exists to rule out non-data user changes.'}
Path('scratch/independent_audit/rerun_verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
