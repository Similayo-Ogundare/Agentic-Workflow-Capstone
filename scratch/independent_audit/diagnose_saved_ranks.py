from pathlib import Path
from collections import defaultdict
import hashlib
import json
import openpyxl

records=[]
source_files=[]
for path in sorted(Path('Marvellous Church/1st Quarter').rglob('*.xlsx')):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    rows=list(wb.active.iter_rows(values_only=True))[1:]
    source_files.append({'path':str(path.resolve()),'members':len(rows),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    for row in rows:
        first,last,tithe,offering=row
        records.append({'name':f'{first} {last}','month':path.parent.name,'tithe':tithe if isinstance(tithe,(int,float)) else 0,'offering':offering if isinstance(offering,(int,float)) else 0,'tithe_paid':int(isinstance(tithe,(int,float)) and tithe>0),'offering_paid':int(isinstance(offering,(int,float)) and offering>0)})
    wb.close()
members={}
for r in records:
    m=members.setdefault(r['name'],{k:0 for k in ('tithe','offering','tithe_paid','offering_paid')})
    for key in m:
        m[key]+=r[key]
for m in members.values():
    m['combined']=m['tithe']+m['offering']

rankings={}
tie_findings={}
for hierarchy in ('A_category_then_combined','B_combined'):
    rankings[hierarchy]={}
    tie_findings[hierarchy]={}
    for kind in ('tithe','offering'):
        key=lambda n:(members[n][kind+'_paid'],members[n][kind],members[n]['combined']) if hierarchy.startswith('A') else (members[n][kind+'_paid'],members[n]['combined'])
        grouped=defaultdict(list)
        for name in members:
            grouped[key(name)].append(name)
        tie_findings[hierarchy][kind]=[{'criteria':criteria,'names':names} for criteria,names in grouped.items() if len(names)>1]
        rankings[hierarchy][kind]={name:1+sum(key(other)>key(name) for other in members) for name in members}

path=Path('outputs/Final_Summary.xlsx')
formula_book=openpyxl.load_workbook(path,read_only=True,data_only=False)
value_book=openpyxl.load_workbook(path,read_only=True,data_only=True)
formula_rows=list(formula_book['Summary'].iter_rows(values_only=True))
value_rows=list(value_book['Summary'].iter_rows(values_only=True))
actual=[]
mismatches=[]
differences=[]
for row_no in range(7,26):
    formula_row=formula_rows[row_no-1]
    row=value_rows[row_no-1]
    name=f'{row[0]} {row[1]}'
    member=members[name]
    item={'name':name,'row':row_no,'tithe':row[2],'offering':row[3],'tithe_paid':row[4],'offering_paid':row[6],'tithe_rank':row[5],'offering_rank':row[7],'tithe_rank_formula':formula_row[5],'offering_rank_formula':formula_row[7],'tithe_rank_A':rankings['A_category_then_combined']['tithe'][name],'offering_rank_A':rankings['A_category_then_combined']['offering'][name],'tithe_rank_B':rankings['B_combined']['tithe'][name],'offering_rank_B':rankings['B_combined']['offering'][name],'combined':member['combined']}
    actual.append(item)
    for col,key in ((2,'tithe'),(3,'offering'),(4,'tithe_paid'),(6,'offering_paid')):
        if row[col]!=member[key]:
            mismatches.append({'name':name,'key':key,'saved':row[col],'source':member[key]})
    for kind in ('tithe','offering'):
        a=rankings['A_category_then_combined'][kind][name]
        b=rankings['B_combined'][kind][name]
        if a!=b:
            differences.append({'name':name,'category':kind,'paid':member[kind+'_paid'],'category_amount':member[kind],'combined':member['combined'],'rank_A':a,'rank_B':b,'delta':b-a})

duplicates={}
for kind in ('tithe','offering'):
    values=[r[kind+'_rank'] for r in actual]
    duplicates[kind]={'values':values,'unique_1_to_19':sorted(values)==list(range(1,20)),'mismatch_A':[r['name'] for r in actual if r[kind+'_rank']!=r[kind+'_rank_A']],'mismatch_B':[r['name'] for r in actual if r[kind+'_rank']!=r[kind+'_rank_B']]}
headers=list(value_rows[5])
saved_explanation=value_rows[28][0]
formula_book.close();value_book.close()

output_paths=[]
for workbook_path in sorted(Path('outputs').rglob('*.xlsx')):
    if workbook_path.name.startswith('~$'):
        output_paths.append({'path':str(workbook_path.resolve()),'excel_lock_file':True})
        continue
    wb=openpyxl.load_workbook(workbook_path,read_only=True,data_only=True)
    rank_headers=[]
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for v in row:
                if isinstance(v,str) and ('Consistency Rank' in v or 'consistency rank' in v.lower()):
                    rank_headers.append({'sheet':sheet.title,'label':v})
    output_paths.append({'path':str(workbook_path.resolve()),'sheets':wb.sheetnames,'consistency_rank_headers':rank_headers,'superseded_Q1':workbook_path.name=='Marvellous_Church_Q1_2025_Contributions.xlsx'})
    wb.close()

result={'instructions_sha256':hashlib.sha256(Path('INSTRUCTIONS.md').read_bytes()).hexdigest(),'actual_final_path':str(path.resolve()),'actual_final_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'source_files':source_files,'source_totals':{k:sum(m[k] for m in members.values()) for k in ('tithe','offering','tithe_paid','offering_paid','combined')},'rank_headers':headers,'saved_explanation_A29':saved_explanation,'actual_saved_members':actual,'source_metric_mismatches':mismatches,'rank_checks':duplicates,'ranking_differences':differences,'unresolved_ties':tie_findings,'output_paths':output_paths}
Path('scratch/independent_audit/saved_rank_diagnostic.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ('source_files','actual_saved_members')},indent=2))
print(json.dumps({'representative_formulas':actual[0]},indent=2))
